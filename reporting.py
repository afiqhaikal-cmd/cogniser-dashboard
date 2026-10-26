"""Factory-format monthly inspection and auditable report exports."""
from io import BytesIO
from collections import Counter
from datetime import datetime
import pandas as pd
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from reportlab.platypus import SimpleDocTemplate,Paragraph,Table,TableStyle,Spacer,PageBreak
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.pagesizes import landscape,A4
from core import connect

SOP = [
('SD-PD-01-01 §1–4','Clean hair, nails, clothing and footwear; wear required protective equipment; no jewellery in production.'),
('SD-PD-01-01 §7','Wash and dry hands and apply approved sanitiser before entering production.'),
('SD-PD-01-01 §16','Report illness or wounds and follow factory-approved protective measures.'),
('SD-PD-01-01 §19','Keep mobile phones and personal items in designated lockers.'),
]
def monthly(month):
 if len(month)!=7 or month[4]!='-':raise ValueError('Month must be YYYY-MM')
 with connect() as c:
  df=pd.read_sql_query('SELECT id AS attempt_id, employee_id, started, finished, stage, status, reason, source FROM attempts WHERE substr(started,1,7)=? ORDER BY id',c,params=(month,))
  checks=pd.read_sql_query('SELECT c.attempt_id,c.key,c.status,c.source,c.time FROM checks c JOIN attempts a ON a.id=c.attempt_id WHERE substr(a.started,1,7)=?',c,params=(month,))
 return df,checks

def summary(df):
 cnt=Counter(df.status.tolist()) if not df.empty else {}
 reasons=Counter(str(x) for x in df.loc[df.status=='FAIL','reason'].dropna() if str(x).strip()) if not df.empty else {}
 return {'attempts':len(df),'pass':cnt.get('PASS',0),'fail':cnt.get('FAIL',0),'pending':cnt.get('PENDING',0),'interrupted':cnt.get('INTERRUPTED',0),'employees':df.employee_id.nunique() if not df.empty else 0,'reasons':reasons}

def factory_pdf(month, df):
 """Factory-reference landscape form: seven original requirements in one left cell,
 days split 1–16 and 17–31; two-page digital adaptation for QA review."""
 from calendar import monthrange
 from reportlab.lib.styles import ParagraphStyle
 from reportlab.lib.enums import TA_CENTER
 from reportlab.pdfbase import pdfmetrics
 from reportlab.pdfbase.ttfonts import TTFont
 from reportlab.platypus import KeepTogether
 y,m=map(int,month.split('-')); nd=monthrange(y,m)[1]
 latest={}
 if not df.empty:
  for _,r in df.iterrows():
   if r['status'] in ('PASS','FAIL'):
    latest[(str(r['employee_id']),str(r['started'])[:10])]=r['status']
 employees=sorted({str(x) for x in df.employee_id}) if not df.empty else []
 font='Helvetica';symbol='Helvetica'
 for path in ('C:/Windows/Fonts/arial.ttf','/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'):
  if __import__('os').path.exists(path):
   try:pdfmetrics.registerFont(TTFont('FactoryUnicode',path));font=symbol='FactoryUnicode';break
   except Exception:pass
 buf=BytesIO()
 doc=SimpleDocTemplate(buf,pagesize=landscape(A4),leftMargin=23,rightMargin=23,topMargin=20,bottomMargin=20)
 normal=ParagraphStyle('factory',fontName=font,fontSize=7.3,leading=10)
 small=ParagraphStyle('factorysmall',parent=normal,fontSize=6.2,leading=8)
 title=ParagraphStyle('factorytitle',parent=normal,fontSize=12,leading=16,alignment=TA_CENTER,spaceAfter=6)
 mid=ParagraphStyle('factorymid',parent=normal,fontSize=9,leading=12,alignment=TA_CENTER)
 items=['1. Employees are properly dressed in accordance with regulations, wear no jewelry, and wear the required footwear.',
 '2. Employees are free from communicable and skin diseases.',
 '3. Fingernails must be kept short.',
 '4. Wounds must be covered with a medical plaster and an additional glove must be worn over it.',
 '5. The face mask must be properly worn.',
 '6. The hairnet must properly cover the hair.',
 '7. The prescribed factory work uniform must be worn.']
 requirements=Paragraph('<br/>'.join(items)+'<br/><br/><b>Remarks</b><br/>√ = Pass<br/>X = Fail',small)
 story=[]
 for page,(first,last) in enumerate(((1,16),(17,31))):
  if page:story.append(PageBreak())
  story.append(Paragraph('Company ...........................................................',mid))
  story.append(Paragraph('<b>Employee Health and Cleanliness Inspection Record Form</b>',title))
  story.append(Paragraph(f'For the month of {datetime(y,m,1).strftime("%B")}  C.E. {y}  (Days {first}–{last})',mid))
  story.append(Spacer(1,9))
  days=list(range(first,last+1));head=[Paragraph('Health<br/>Inspection Items',small),'No.',Paragraph('Employee<br/>Names',small)]+[str(d) for d in days]+['Remarks']
  rows=[head]
  display=employees or ['']
  # One result cell per employee/day; rows retain full-height merged reference column.
  for idx,employee in enumerate(display,1):
   row=[requirements if idx==1 else '',str(idx) if employee else '',employee or '']
   for day in days:
    status=latest.get((employee,f'{month}-{day:02d}')) if day<=nd else None
    row.append('√' if status=='PASS' else 'X' if status=='FAIL' else '')
   row.append('');rows.append(row)
  # Include blank rows for handwriting / expansion.
  for idx in range(len(display)+1,max(8,len(display)+1)):
   rows.append(['',str(idx),'']+['']*len(days)+[''])
  widths=[150,26,105]+[25]*len(days)+[90]
  t=Table(rows,colWidths=widths,rowHeights=[44]+[38]*(len(rows)-1),repeatRows=1)
  t.setStyle(TableStyle([
   ('GRID',(0,0),(-1,-1),0.5,colors.black),('SPAN',(0,1),(0,-1)),
   ('VALIGN',(0,0),(-1,-1),'MIDDLE'),('VALIGN',(0,1),(0,-1),'TOP'),
   ('FONTNAME',(0,0),(-1,-1),font),('FONTSIZE',(0,0),(-1,-1),7),
   ('ALIGN',(1,0),(-1,-1),'CENTER'),('ALIGN',(2,1),(2,-1),'LEFT'),
   ('LEFTPADDING',(0,0),(-1,-1),3),('RIGHTPADDING',(0,0),(-1,-1),3),
   ('TOPPADDING',(0,0),(-1,-1),4),('BOTTOMPADDING',(0,0),(-1,-1),4)]))
  story.append(t);story.append(Spacer(1,15))
  story.append(Paragraph('Inspector ........................................ (Date ....../....../......)     Reviewer ........................................ (Date ....../....../......)',normal))
  story.append(Spacer(1,8))
  story.append(Paragraph('Factory Manager                                      Team Leader FSTL / QCTL',normal))
  story.append(Spacer(1,15))
  story.append(Paragraph('Rev : 0                                             Effective Date : 16 December 2025',mid))
  story.append(Paragraph('FM-PD-01-01',mid))
 doc.build(story)
 return buf.getvalue()

def word_report(month,df,checks,analysis=''):
 stat=summary(df);d=Document();sec=d.sections[0];sec.header.paragraphs[0].text='FOOD SAFETY MONITORING-ISO22000  |  HYGIENE INSPECTION'
 normal=d.styles['Normal'];normal.font.name='Aptos';normal.font.size=Pt(9)
 d.add_heading('Monthly Hygiene Inspection Report',0);d.add_paragraph(f'Report period: {month}  |  Form reference: FM-PD-01-01  |  SOP: SD-PD-01-01')
 d.add_paragraph('Internal inspection record for QA review. Final compliance assessment requires authorized verification.')
 d.add_heading('1. Executive Summary',level=1)
 t=d.add_table(rows=0,cols=2);t.style='Light Shading Accent 1'
 for k,v in [('Total attempts',stat['attempts']),('PASS',stat['pass']),('FAIL',stat['fail']),('PENDING',stat['pending']),('INTERRUPTED',stat['interrupted']),('Unique employees',stat['employees'])]:
  cells=t.add_row().cells;cells[0].text=k;cells[1].text=str(v)
 d.add_heading('2. Inspection Attempts',level=1)
 cols=['Attempt','Employee ID','Date / Time','Temperature','PPE / Check Evidence','Final Status','Reason']
 t=d.add_table(rows=1,cols=len(cols));t.style='Table Grid'
 for i,k in enumerate(cols):t.rows[0].cells[i].text=k
 checkmap={}
 if not checks.empty:
  for _,r in checks.iterrows():checkmap.setdefault(r['attempt_id'],[]).append(f"{r['key']}: {r['status']}")
 if df.empty:
  cells=t.add_row().cells;cells[0].text='No attempts recorded'
 else:
  for _,r in df.iterrows():
   cells=t.add_row().cells;details=checkmap.get(r['attempt_id'],[])
   temp=next((x for x in details if x.startswith('temperature:')),'Not recorded')
   vals=[str(r['attempt_id']),str(r['employee_id']),str(r['started']),temp,', '.join(x for x in details if not x.startswith('temperature:')) or 'Not recorded',str(r['status']),str(r['reason'] or '—')]
   for c,v in zip(cells,vals):c.text=v
 for row in t.rows:
  for cell in row.cells:
   for p in cell.paragraphs:
    for run in p.runs:run.font.size=Pt(7)
 d.add_heading('3. Failure Reasons and Corrective Actions',level=1)
 if stat['reasons']:
  for reason,n in stat['reasons'].most_common():d.add_paragraph(f'{reason}: {n} recorded FAIL attempt(s). Corrective action requires operator/QA confirmation.',style='List Bullet')
 else:d.add_paragraph('No recorded failures for this month. Pending attempts are not counted as failures.')
 d.add_heading('4. Factory SOP Reference',level=1)
 for ref,rule in SOP:d.add_paragraph(f'{ref}: {rule}',style='List Bullet')
 d.add_heading('5. Compliance Review',level=1)
 d.add_paragraph(analysis.strip() if analysis.strip() else 'No supplementary compliance commentary recorded. Inspection statistics are based on recorded data.')
 d.add_paragraph('Prepared by: ____________________    Reviewed by: ____________________')
 d.add_paragraph('QA review and approval required before official use.')
 b=BytesIO();d.save(b);return b.getvalue()

def ollama_report(month,df,model='llama3.2:1b'):
 import urllib.request,json,urllib.error
 stat=summary(df)
 data={'month':month,'summary':{k:v for k,v in stat.items() if k!='reasons'},'failure_reasons':dict(stat['reasons'])}
 prompt=(
  "Prepare a concise, evidence-grounded monthly food safety and personnel hygiene compliance review in English. "
  "Use the headings: Recorded Findings, Outstanding Verification, Corrective Action Recommendations, ISO Reference, QA Review. "
  "Use ONLY the provided recorded audit counts and failure reasons; do not invent events, dates, actions, causes, health conditions or approvals. "
  "A PENDING or INTERRUPTED inspection is incomplete, never automatically FAIL or PASS. "
  "Do not infer an ISO nonconformity from a detection alone. Explain what additional QA evidence is required. "
  "Cite ISO 22002-100:2025 Clause 14 as a contextual PRP reference only when relevant: "
  "14.1 General; 14.2 Hygiene facilities; 14.3 Designated eating areas; "
  "14.4 Workwear and protective clothing; 14.5 Health status; "
  "14.6 Personnel cleanliness; 14.7 Personnel behaviour; 14.8 Visitors and external providers. "
  "Do not quote or fabricate ISO requirements or assert that certification has been achieved. "
  "Factory SOP SD-PD-01-01 is separate from ISO and must be verified against the approved factory copy. "
  "Never assume CCTV confirms handwashing, body temperature, illness or jewellery absence without validated evidence. "
  "Do not recommend automatic gate authorization based on these records. "
  "If data are insufficient, explicitly state Not established from available records. "
  "FACTORY SOP REFERENCE: " + json.dumps(SOP,ensure_ascii=False) +
  "\nRECORDED AUDIT DATA: " + json.dumps(data,ensure_ascii=False)
 )
 body=json.dumps({'model':model,'prompt':prompt,'stream':False}).encode()
 req=urllib.request.Request('http://127.0.0.1:11434/api/generate',data=body,headers={'Content-Type':'application/json'})
 with urllib.request.urlopen(req,timeout=150) as response:return json.loads(response.read())['response']
